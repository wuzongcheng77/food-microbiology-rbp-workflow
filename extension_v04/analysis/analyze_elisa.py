from pathlib import Path
import json,csv
import numpy as np
from scipy.optimize import curve_fit
from scipy.stats import ttest_rel
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'04_ELISA'
def savecsv(name,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def model(x,b,k):return b*x/(k+x)
w=load_workbook(ROOT/'01_source_snapshot/ELISA/结合活性.xlsx',data_only=True).active
x=np.array([200,150,125,100,75,50,30,20,10,5,2.5,1.])
y=np.array([[w.cell(i,j).value for j in range(2,14)] for i in range(2,5)],float)
bg=np.mean([w.cell(5,j).value for j in range(2,5)])
z=y-bg
fits=[]
for i in range(3):
    p,c=curve_fit(model,x,z[i],p0=[2,40],maxfev=20000)
    r2=1-np.sum((z[i]-model(x,*p))**2)/np.sum((z[i]-z[i].mean())**2)
    fits.append({'experiment_id':i+1,'Bmax':p[0],'apparent_Kd_ug_mL':p[1],'fit_SE_Kd':np.sqrt(c[1,1]),'R_squared':r2})
savecsv('binding_independent_experiment_fits.csv',fits)
savecsv('binding_experiment_means.csv',[{'experiment_id':i+1,'concentration_ug_mL':v,'stored_mean_OD450':y[i,j],'shared_recorded_background':bg,'corrected_mean_OD450':z[i,j],'technical_wells_averaged_author_confirmed':3} for i in range(3) for j,v in enumerate(x)])
w=load_workbook(ROOT/'01_source_snapshot/ELISA/特异性.xlsx',data_only=True)['特异性']
labels=['O:3','O:8','O:5,27','O:5 LHF8-2','O:5 LHF8-3','O:6,30','L. monocytogenes','V. parahaemolyticus','E. coli']
raw=np.array([[w.cell(i,j).value for j in range(2,11)] for i in range(2,5)],float)
bgs=np.array([np.mean([w.cell(i,j).value for i in (5,6)]) for j in range(2,11)])
sp=raw-bgs
tests=[]
for j in range(1,9):
    t,p=ttest_rel(sp[:,0],sp[:,j])
    tests.append({'comparator':labels[j],'mean_target_minus_comparator':(sp[:,0]-sp[:,j]).mean(),'paired_t':t,'df':2,'unadjusted_p':p})
order=np.argsort([v['unadjusted_p'] for v in tests]);last=0
for k,j in enumerate(order):
    last=max(last,min(1,(8-k)*tests[j]['unadjusted_p']));tests[j]['Holm_p']=last
savecsv('specificity_paired_Holm_analysis.csv',tests)
savecsv('specificity_experiment_means.csv',[{'experiment_id_assuming_matched_rows':i+1,'group':labels[j],'stored_mean_OD450':raw[i,j],'group_background':bgs[j],'corrected_mean_OD450':sp[i,j]} for i in range(3) for j in range(9)])
ks=[a['apparent_Kd_ug_mL'] for a in fits]
summary={'independent_n':3,'technical_wells_per_experiment':3,'Kd_mean':float(np.mean(ks)),'Kd_between_experiment_SD':float(np.std(ks,ddof=1)),'Kd_each':ks,'background_binding':bg,'specificity_means':sp.mean(axis=0).tolist(),'specificity_SD':sp.std(axis=0,ddof=1).tolist(),'specificity_paired_tests_status':'Author confirmed matched experiment rows and source background treatment on 2026-09-29. Retrospective two-sided paired t tests with Holm correction across eight comparisons. Small n limits assumption checks.','limitations':['Underlying technical well values are not present in these source sheets; three stored values are experiment means according to author confirmation.','Binding background is shared as stored. Specificity background is group-specific as stored. Author confirmed these treatments. Background measurement uncertainty is not propagated.','New analysis is retrospective. Kd SD measures between-experiment spread, not pooled-fit SE.']}
(OUT/'ELISA_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
fig,ax=plt.subplots(1,2,figsize=(11,4.1),gridspec_kw={'width_ratios':[1,1.3]})
colors=['#2878B5','#E07A3F','#328C68'];xx=np.linspace(0,205,250)
for i,fit in enumerate(fits):
    ax[0].scatter(x,z[i],s=18,color=colors[i],label=f'Experiment {i+1}')
    ax[0].plot(xx,model(xx,fit['Bmax'],fit['apparent_Kd_ug_mL']),color=colors[i],lw=1)
ax[0].set(xlabel='Fusion preparation (µg/mL)',ylabel='Background-corrected OD450',title='A  Independent binding experiments');ax[0].legend(frameon=False,fontsize=8)
pos=np.arange(9);ax[1].bar(pos,sp.mean(axis=0),color='#C8DCE8',edgecolor='#38627A',width=.65)
ax[1].errorbar(pos,sp.mean(axis=0),sp.std(axis=0,ddof=1),fmt='none',color='#333333',capsize=3)
for i in range(3):ax[1].scatter(pos+(i-1)*.14,sp[i],s=18,color=colors[i],zorder=3)
ax[1].set_xticks(pos,labels,rotation=60,ha='right',fontsize=8);ax[1].set(ylabel='Background-corrected OD450',title='B  Tested specificity panel')
for a in ax:a.spines[['top','right']].set_visible(False)
fig.tight_layout();fig.savefig(ROOT/'06_figures/Figure_4_ELISA.pdf');fig.savefig(ROOT/'06_figures/Figure_4_ELISA.png',dpi=220);plt.close(fig)
print(json.dumps(summary,indent=2))
